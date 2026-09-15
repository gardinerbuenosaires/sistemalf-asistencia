"""Prueba de los puestos que no reciben uniforme."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _base import preparar, token_sistema
DB = preparar()   # una COPIA de la base: la real nunca se toca

import io, sqlite3
from datetime import date
import openpyxl
from fastapi.testclient import TestClient
import main
from auth.core import create_token, ensure_admin, invalidar_cache

ensure_admin()
ok = fallos = 0
creadas = []


def chequear(desc, cond, extra=""):
    global ok, fallos
    if cond:
        ok += 1
        print(f"  OK   {desc}")
    else:
        fallos += 1
        print(f"  FALLA {desc}  {extra}")


con = sqlite3.connect(DB)
con.row_factory = sqlite3.Row
con.execute("UPDATE configuracion SET valor='1' WHERE clave='uniformes_activo'")
con.commit()
MARCAS_PREVIAS = [tuple(r) for r in con.execute(
    "SELECT cargo_id, marcado_por, marcado_en FROM uniformes_cargos_sin_uniforme")]
TALLES_PREVIOS = [tuple(r) for r in con.execute(
    "SELECT empleado_id, tipo_talle_id, valor, modificado_en FROM uniformes_talles_empleado")]
con.execute("DELETE FROM uniformes_cargos_sin_uniforme")
con.commit()

cli = TestClient(main.app)
cli.cookies.set("session", token_sistema()[0])

# Un cargo con al menos dos personas que no recibieron nada todavía.
fila = con.execute("""
    SELECT e.cargo_id, group_concat(e.id) ids FROM empleados e
    WHERE e.activo=1 AND e.tipo!='acceso' AND e.cargo_id IS NOT NULL
      AND e.id NOT IN (SELECT empleado_id FROM uniformes_movimientos)
    GROUP BY e.cargo_id HAVING COUNT(*) >= 2 LIMIT 1""").fetchone()
CARGO = fila["cargo_id"]
NUNCA, RECIBIO = [int(x) for x in fila["ids"].split(",")[:2]]
el = next(e for e in cli.get("/api/uniformes/elementos?solo_activos=true").json() if not e["tipo_talle_id"])
r = cli.post("/api/uniformes/constancias", json={
    "empleado_id": RECIBIO, "fecha": date.today().isoformat(),
    "items": [{"elemento_id": el["id"], "cantidad": 1}]})
creadas.append(r.json()["id"])


def ids_reporte(**q):
    qs = "&".join(f"{k}={v}" for k, v in q.items())
    d = cli.get(f"/api/uniformes/reportes/antiguedad?{qs}").json()
    return {f["id"]: f for f in d["filas"]}, d


print("\n=== LA LISTA DE PUESTOS ===")
p = cli.get("/api/uniformes/puestos").json()
cargos = con.execute("SELECT COUNT(*) FROM cargos").fetchone()[0]
chequear("lista todos los cargos", len(p) == cargos, f"{len(p)} vs {cargos}")
chequear("sin marcas, todos reciben", all(x["recibe"] for x in p))
chequear("cada puesto trae su cantidad de empleados", all("empleados" in x for x in p))

print("\n=== MARCAR UN PUESTO ===")
r = cli.put(f"/api/uniformes/puestos/{CARGO}", json={"recibe": False})
chequear("marcar como «no recibe» -> 200", r.status_code == 200, r.status_code)
x = next(x for x in cli.get("/api/uniformes/puestos").json() if x["id"] == CARGO)
chequear("queda marcado, con quién lo marcó", x["recibe"] is False and bool(x["marcado_por"]), x)
chequear("marcar dos veces no falla",
         cli.put(f"/api/uniformes/puestos/{CARGO}", json={"recibe": False}).status_code == 200)

filas, d = ids_reporte()
chequear("quien nunca recibió desaparece del reporte", NUNCA not in filas)
chequear("quien SÍ recibió sigue apareciendo aunque su puesto esté marcado", RECIBIO in filas)
chequear("y aparece marcado como puesto que no recibe", filas.get(RECIBIO, {}).get("sin_uniforme") is True)
chequear("el reporte dice cuántos quedaron ocultos", d["ocultos_sin_uniforme"] >= 1, d["ocultos_sin_uniforme"])

filas, d = ids_reporte(incluir_sin_uniforme="true")
chequear("con la casilla, vuelven a aparecer", NUNCA in filas and filas[NUNCA]["sin_uniforme"] is True)
chequear("y no hay ocultos", d["ocultos_sin_uniforme"] == 0)

det = cli.get(f"/api/uniformes/reportes/entregas?empleado_id={RECIBIO}").json()["filas"]
chequear("el reporte de Entregas no se ve afectado", len(det) == 1)

print("\n=== EXCEL ===")
nombre = con.execute("SELECT apellido || ', ' || nombre FROM empleados WHERE id=?", (NUNCA,)).fetchone()[0]


def nombres_excel(q=""):
    ws = openpyxl.load_workbook(io.BytesIO(cli.get(f"/api/uniformes/reportes/antiguedad.xlsx{q}").content)).active
    return [ws.cell(i, 1).value for i in range(5, ws.max_row + 1)], ws["A2"].value


n, sub = nombres_excel()
chequear("el Excel respeta la marca", nombre not in n)
chequear("y lo aclara en el subtítulo", "no reciben uniforme" in (sub or ""), sub)
n, _ = nombres_excel("?incluir_sin_uniforme=true")
chequear("con la casilla, el Excel los incluye", nombre in n)

print("\n=== DESMARCAR ===")
cli.put(f"/api/uniformes/puestos/{CARGO}", json={"recibe": True})
filas, _ = ids_reporte()
chequear("al volver a tildarlo, reaparece quien nunca recibió", NUNCA in filas)

print("\n=== VALIDACIONES Y PERMISOS ===")
chequear("un cargo inexistente -> 404",
         cli.put("/api/uniformes/puestos/999999", json={"recibe": False}).status_code == 404)
rol = con.execute("""SELECT rol_id FROM permisos WHERE modulo='uniformes' AND accion='ver'
                     AND rol_id NOT IN (SELECT rol_id FROM permisos WHERE modulo='uniformes' AND accion='editar')
                     LIMIT 1""").fetchone()
if rol:
    solo_ver = TestClient(main.app)
    solo_ver.cookies.set("session", create_token(999, "prueba@local", rol["rol_id"], "prueba"))
    chequear("quien solo ve, puede consultar la lista", solo_ver.get("/api/uniformes/puestos").status_code == 200)
    chequear("pero no puede cambiarla (hace falta uniformes:editar)",
             solo_ver.put(f"/api/uniformes/puestos/{CARGO}", json={"recibe": False}).status_code == 403)
con.execute("UPDATE configuracion SET valor='0' WHERE clave='uniformes_activo'")
con.commit()
chequear("con la bandera en 0, responde 404", cli.get("/api/uniformes/puestos").status_code == 404)
con.execute("UPDATE configuracion SET valor='1' WHERE clave='uniformes_activo'")
con.commit()

# ── Limpieza ────────────────────────────────────────────────────────────────
for cid in creadas:
    con.execute("DELETE FROM uniformes_items WHERE movimiento_id=?", (cid,))
    con.execute("DELETE FROM uniformes_movimientos WHERE id=?", (cid,))
con.execute("DELETE FROM uniformes_cargos_sin_uniforme")
con.executemany("INSERT INTO uniformes_cargos_sin_uniforme (cargo_id, marcado_por, marcado_en) VALUES (?,?,?)",
                MARCAS_PREVIAS)
con.execute("DELETE FROM uniformes_talles_empleado")
con.executemany("INSERT INTO uniformes_talles_empleado (empleado_id, tipo_talle_id, valor, modificado_en) "
                "VALUES (?,?,?,?)", TALLES_PREVIOS)
con.commit()
con.close()
invalidar_cache()

print(f"\n{'='*52}\n  {ok} pasaron, {fallos} fallaron\n{'='*52}")
raise SystemExit(1 if fallos else 0)
