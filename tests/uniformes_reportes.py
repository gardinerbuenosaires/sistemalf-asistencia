"""Prueba de la tanda 5a: reportes.

Arma entregas con fechas y rubros conocidos sobre empleados que no tienen
ninguna, y compara contra lo que había antes: así las constancias reales que
ya existan en la base no confunden las cuentas.
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _base import preparar, token_sistema
DB = preparar()   # una COPIA de la base: la real nunca se toca

import io, sqlite3
from datetime import date
from fastapi.testclient import TestClient
import main
from auth.core import create_token, ensure_admin, invalidar_cache

ensure_admin()
ok = fallos = 0
creadas = []
HOY = date.today().isoformat()


def chequear(desc, cond, extra=""):
    global ok, fallos
    if cond:
        ok += 1
        print(f"  OK   {desc}")
    else:
        fallos += 1
        print(f"  FALLA {desc}  {extra}")


con = sqlite3.connect(DB)
con.row_factory = sqlite3.Row
con.execute("UPDATE configuracion SET valor='1' WHERE clave='uniformes_activo'")
con.commit()
TALLES_PREVIOS = [tuple(r) for r in con.execute(
    "SELECT empleado_id, tipo_talle_id, valor, modificado_en FROM uniformes_talles_empleado")]

cli = TestClient(main.app)
cli.cookies.set("session", token_sistema()[0])

# ── Datos de prueba ──────────────────────────────────────────────────────────
libres = [r["id"] for r in con.execute("""
    SELECT e.id FROM empleados e
    WHERE e.activo=1 AND e.tipo!='acceso' AND e.cargo_id IS NOT NULL
      AND e.id NOT IN (SELECT empleado_id FROM uniformes_movimientos)
    ORDER BY e.id LIMIT 4""")]
A, B, C, D = libres
rubros = {r["nombre"]: r for r in cli.get("/api/uniformes/rubros").json()}
ropa, epp = rubros["Ropa de trabajo"], rubros["EPP"]
els = cli.get("/api/uniformes/elementos?solo_activos=true").json()
el_ropa = next(e for e in els if e["categoria_id"] == ropa["id"] and e["tipo_talle_id"])
el_epp = next(e for e in els if e["categoria_id"] == epp["id"] and not e["tipo_talle_id"])
escala = [v["valor"] for v in cli.get(f"/api/uniformes/tipos-talle/{el_ropa['tipo_talle_id']}/valores").json()]
t1, t2 = escala[0], escala[1]

base_resumen = {f["elemento_id"]: f for f in
                cli.get("/api/uniformes/reportes/entregas?vista=resumen").json()["filas"]}


def emitir(emp, fecha, items, **extra):
    r = cli.post("/api/uniformes/constancias", json={"empleado_id": emp, "fecha": fecha, "items": items, **extra})
    assert r.status_code == 201, r.text
    creadas.append(r.json()["id"])
    return r.json()


emitir(A, "2024-01-10", [{"elemento_id": el_ropa["id"], "cantidad": 2, "talle": t1}])
emitir(A, HOY, [{"elemento_id": el_epp["id"], "cantidad": 1}])
emitir(B, HOY, [{"elemento_id": el_ropa["id"], "cantidad": 1, "talle": t2}])
emitir(C, "2023-06-01", [{"elemento_id": el_ropa["id"], "cantidad": 1, "talle": t1}], origen="historico")
anulada = emitir(B, HOY, [{"elemento_id": el_epp["id"], "cantidad": 5}])
cli.post(f"/api/uniformes/constancias/{anulada['id']}/anular", json={"motivo": "prueba"})
emitir(B, HOY, [{"elemento_id": el_epp["id"], "cantidad": 7}], tipo="devolucion")

print("\n=== ENTREGAS — DETALLE ===")
d = cli.get(f"/api/uniformes/reportes/entregas?empleado_id={A}").json()
chequear("A tiene sus dos renglones", len(d["filas"]) == 2, len(d["filas"]))
chequear("ordenado del más nuevo al más viejo", [f["fecha"] for f in d["filas"]] == [HOY, "2024-01-10"])
d = cli.get(f"/api/uniformes/reportes/entregas?empleado_id={B}").json()
chequear("de B no cuentan la anulada ni la devolución", [f["cantidad"] for f in d["filas"]] == [1],
         [f["cantidad"] for f in d["filas"]])
d = cli.get(f"/api/uniformes/reportes/entregas?empleado_id={C}").json()
chequear("la histórica del papel sí cuenta", len(d["filas"]) == 1 and d["filas"][0]["numero"] is None)
d = cli.get(f"/api/uniformes/reportes/entregas?empleado_id={A}&hasta=2024-12-31").json()
chequear("filtro por período", [f["fecha"] for f in d["filas"]] == ["2024-01-10"])
d = cli.get(f"/api/uniformes/reportes/entregas?empleado_id={A}&rubro_id={epp['id']}").json()
chequear("filtro por rubro", len(d["filas"]) == 1 and d["filas"][0]["rubro"] == "EPP")
cargo_a = con.execute("SELECT cargo_id FROM empleados WHERE id=?", (A,)).fetchone()[0]
d = cli.get(f"/api/uniformes/reportes/entregas?cargo_id={cargo_a}").json()
ids = {r["empleado_id"] for r in con.execute(
    f"SELECT empleado_id FROM uniformes_movimientos WHERE id IN ({','.join('?'*len(creadas))})", creadas)}
chequear("filtro por cargo: todas las filas son de ese cargo",
         d["filas"] and all(f["cargo"] == con.execute("SELECT nombre FROM cargos WHERE id=?", (cargo_a,)).fetchone()[0]
                            for f in d["filas"]))
chequear("vista inválida -> 400", cli.get("/api/uniformes/reportes/entregas?vista=zz").status_code == 400)

print("\n=== ENTREGAS — RESUMEN ===")
res = {f["elemento_id"]: f for f in cli.get("/api/uniformes/reportes/entregas?vista=resumen").json()["filas"]}
b = base_resumen.get(el_ropa["id"], {"unidades": 0, "constancias": 0, "personas": 0, "talles": []})
r = res[el_ropa["id"]]
chequear("unidades del elemento de ropa: +4 (2 de A, 1 de B, 1 histórica de C)",
         r["unidades"] - b["unidades"] == 4, r["unidades"] - b["unidades"])
chequear("constancias: +3", r["constancias"] - b["constancias"] == 3)
chequear("personas: +3", r["personas"] - b["personas"] == 3)
bt = {x["talle"]: x["unidades"] for x in b["talles"]}
rt = {x["talle"]: x["unidades"] for x in r["talles"]}
chequear(f"por talle: {t1} +3 y {t2} +1",
         rt.get(t1, 0) - bt.get(t1, 0) == 3 and rt.get(t2, 0) - bt.get(t2, 0) == 1, rt)
be = base_resumen.get(el_epp["id"], {"unidades": 0})
chequear("EPP: +1 (la anulada de 5 y la devolución de 7 no suman)",
         res[el_epp["id"]]["unidades"] - be["unidades"] == 1, res[el_epp["id"]]["unidades"] - be["unidades"])

print("\n=== ÚLTIMA ENTREGA POR PERSONA ===")
a = cli.get("/api/uniformes/reportes/antiguedad").json()
pos = {f["id"]: i for i, f in enumerate(a["filas"])}
fila = {f["id"]: f for f in a["filas"]}
chequear("aparece quien nunca recibió nada (D)", D in fila and fila[D]["ultima"] is None)
chequear("los que nunca recibieron van antes que todos los que sí",
         max(pos[f["id"]] for f in a["filas"] if f["ultima"] is None)
         < min(pos[f["id"]] for f in a["filas"] if f["ultima"] is not None))
fr = fila[A]["rubros"][str(ropa["id"])]
chequear("A: la última ropa es la de 2024-01-10", fr and fr["fecha"] == "2024-01-10")
esperada = ropa["meses_alerta"] is not None and fr["meses"] >= ropa["meses_alerta"]
chequear(f"A: el aviso de ropa respeta el umbral del rubro ({ropa['meses_alerta']} meses)",
         fr["alerta"] == esperada, fr)
fe = fila[A]["rubros"][str(epp["id"])]
chequear("A: EPP de hoy, sin aviso", fe and fe["fecha"] == HOY and fe["alerta"] is False)
chequear("A: la última de cualquier cosa es la de hoy", fila[A]["ultima"]["fecha"] == HOY)
chequear("B: la anulada y la devolución no cuentan como EPP", fila[B]["rubros"][str(epp["id"])] is None)
chequear("C: la histórica del papel cuenta", fila[C]["rubros"][str(ropa["id"])]["fecha"] == "2023-06-01")
o = cli.get(f"/api/uniformes/reportes/antiguedad?rubro_id={ropa['id']}").json()
orden = [f["id"] for f in o["filas"] if f["id"] in (A, B, C, D)]
chequear("ordenado por ropa: nunca, después de la más vieja a la más nueva (D, C, A, B)",
         orden == [D, C, A, B], orden)

print("\n=== EXCEL ===")
import openpyxl
for url, titulo in [("/api/uniformes/reportes/entregas.xlsx?vista=detalle", "Entregas — detalle"),
                    ("/api/uniformes/reportes/entregas.xlsx?vista=resumen", "Entregas — resumen"),
                    ("/api/uniformes/reportes/antiguedad.xlsx", "Última entrega por persona")]:
    r = cli.get(url)
    tipo_ok = r.headers.get("content-type", "").startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    try:
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        abre = ws["A1"].value == titulo
    except Exception as e:
        abre = False
    chequear(f"{titulo}: descarga un Excel válido", r.status_code == 200 and tipo_ok and abre,
             f"{r.status_code} {r.headers.get('content-type')}")
json_det = cli.get("/api/uniformes/reportes/entregas?vista=detalle").json()["filas"]
ws = openpyxl.load_workbook(io.BytesIO(cli.get("/api/uniformes/reportes/entregas.xlsx?vista=detalle").content)).active
chequear("el Excel del detalle tiene las mismas filas que la pantalla", ws.max_row - 4 == len(json_det),
         f"{ws.max_row - 4} vs {len(json_det)}")
chequear("la fecha va como fecha de Excel, no como texto",
         ws.cell(5, 1).is_date and ws.cell(5, 1).number_format == "DD/MM/YYYY")

print("\n=== PERMISOS Y BANDERA ===")
sin = con.execute("""SELECT r.id FROM roles r WHERE r.id NOT IN
                     (SELECT rol_id FROM permisos WHERE modulo='uniformes' AND accion='ver') LIMIT 1""").fetchone()
if sin:
    otro = TestClient(main.app)
    otro.cookies.set("session", create_token(999, "prueba@local", sin["id"], "prueba"))
    chequear("un rol sin uniformes:ver no puede ver los reportes",
             otro.get("/api/uniformes/reportes/entregas").status_code == 403)
con.execute("UPDATE configuracion SET valor='0' WHERE clave='uniformes_activo'")
con.commit()
chequear("con la bandera en 0, los reportes responden 404",
         cli.get("/api/uniformes/reportes/antiguedad").status_code == 404)
con.execute("UPDATE configuracion SET valor='1' WHERE clave='uniformes_activo'")
con.commit()

# ── Limpieza ────────────────────────────────────────────────────────────────
for cid in creadas:
    con.execute("DELETE FROM uniformes_items WHERE movimiento_id=?", (cid,))
    con.execute("DELETE FROM uniformes_movimientos WHERE id=?", (cid,))
con.execute("DELETE FROM uniformes_talles_empleado")
con.executemany("INSERT INTO uniformes_talles_empleado (empleado_id, tipo_talle_id, valor, modificado_en) "
                "VALUES (?,?,?,?)", TALLES_PREVIOS)
con.commit()
con.close()
invalidar_cache()

print(f"\n{'='*52}\n  {ok} pasaron, {fallos} fallaron\n{'='*52}")
raise SystemExit(1 if fallos else 0)
